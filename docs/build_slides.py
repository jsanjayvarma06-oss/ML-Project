"""Build the review slide deck (16:9 .pptx) from results/metrics.json and figures.

    python docs/build_slides.py
"""
import json
import sys
from pathlib import Path

import numpy as np
from PIL import Image as PILImage
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from linkfail import config as C  # noqa: E402
from linkfail.simulator import Fault, LinkDesign, snapshot  # noqa: E402

OUT = ROOT / "docs" / "Slides_Link_Failure_Prediction.pptx"
ASSETS = ROOT / "docs" / "assets"

SANJAY, KANAK = "J Sanjay Varma", "Kanak Pandey"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
AQUA = RGBColor(0x1B, 0xAF, 0x7A)
INK = RGBColor(0x1A, 0x1A, 0x19)
MUTED = RGBColor(0x52, 0x51, 0x4E)
GRID = RGBColor(0xE4, 0xE3, 0xDF)
GREY = RGBColor(0xB0, 0xAF, 0xA8)
TINT = RGBColor(0xEE, 0xF3, 0xFA)
TINT_O = RGBColor(0xFD, 0xEF, 0xE9)
TINT_A = RGBColor(0xE6, 0xF6, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
HEAD, BODY, MONO = "Cambria", "Calibri", "Courier New"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
_count = [0]


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------
def _runs(p, parts, size, color, font, bold):
    for t, o in parts:
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = o.get("font", font)
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.italic = o.get("italic", False)
        f.color.rgb = o.get("color", color)


def text(slide, x, y, w, h, runs, size=16, color=INK, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6):
    """runs: str | list of paragraphs; a paragraph is str or list of (text, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        _runs(p, [(para, {})] if isinstance(para, str) else para, size, color, font, bold)
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=9, marker_color=ORANGE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        _runs(p, [("▸  ", {"color": marker_color})], size, color, BODY, False)
        _runs(p, [(item, {})] if isinstance(item, str) else item, size, color, BODY, False)
    return tb


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.25, radius=None, rot=0):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if radius is not None:
        s.adjustments[0] = radius
    s.rotation = rot
    s.shadow.inherit = False
    return s


def card(slide, x, y, w, h, fill=TINT, radius=0.06):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, radius=radius)


def label_in(s, t, size=14, color=WHITE, bold=True, font=BODY):
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _runs(p, [(t, {})], size, color, font, bold)


def badge(slide, x, y, n, d=0.55, fill=ORANGE, size=16):
    c = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=fill)
    label_in(c, str(n), size=size)
    return c


def line(slide, x1, y1, x2, y2, color=GREY, w=2, dash=None):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(w)
    if dash:
        ln.line.dash_style = dash
    return ln


def arrow(slide, x, y, w=0.32, h=0.3, color=ORANGE):
    return shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, fill=color)


def picture(slide, path, x, y, w=None, h=None, border=False):
    path = Path(path)
    iw, ih = PILImage.open(path).size
    if w and h:                               # fit inside box, keep aspect, centre
        if iw / ih > w / h:
            h2 = w * ih / iw
            y, h = y + (h - h2) / 2, h2
        else:
            w2 = h * iw / ih
            x, w = x + (w - w2) / 2, w2
    elif w:
        h = w * ih / iw
    else:
        w = h * iw / ih
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if border:
        pic.line.color.rgb = GRID
        pic.line.width = Pt(1)
    return pic


def figure(slide, name, *a, **k):
    return picture(slide, C.FIGURES / f"{name}.png", *a, **k)


def stat(slide, x, y, w, value, label, color=BLUE, fill=TINT, vsize=32, h=1.35):
    card(slide, x, y, w, h, fill)
    text(slide, x + 0.22, y + 0.13, w - 0.44, 0.65, value, size=vsize, bold=True, color=color, font=HEAD)
    text(slide, x + 0.22, y + 0.82, w - 0.44, h - 0.9, label, size=12, color=MUTED)


def notes(slide, presenter, body):
    slide.notes_slide.notes_text_frame.text = f"Presenter: {presenter}\n\n{body}"


def new_slide(section=None, presenter=None, dark=False):
    s = prs.slides.add_slide(BLANK)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY if dark else WHITE
    _count[0] += 1
    n = _count[0]
    if section:                                   # section chip, top right
        chip = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 11.05, 0.5, 1.68, 0.36,
                     fill=RGBColor(0x16, 0x33, 0x5C) if dark else TINT, radius=0.5)
        label_in(chip, section.upper(), size=10, color=ICE if dark else BLUE)
    if n > 1:
        text(s, 12.0, 7.02, 0.73, 0.3, str(n), size=11, color=ICE if dark else MUTED, align=PP_ALIGN.RIGHT)
    if presenter:
        text(s, 0.6, 7.02, 6, 0.3, f"Presenter · {presenter}", size=10, color=ICE if dark else GREY)
    return s


def title(slide, t, sub=None, dark=False, w=10.2):
    text(slide, 0.6, 0.42, w, 0.8, t, size=32, font=HEAD, bold=True, color=WHITE if dark else NAVY)
    if sub:
        text(slide, 0.6, 1.14, 11.5, 0.45, sub, size=16, color=ICE if dark else MUTED)


def style_chart(ch, legend=True, size=12):
    ch.font.name = BODY
    ch.font.size = Pt(size)
    ch.font.color.rgb = MUTED
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)
        ch.legend.font.color.rgb = INK
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = GRID
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(size - 1)
    ch.category_axis.format.line.color.rgb = GREY
    ch.category_axis.tick_labels.font.size = Pt(size)


def axis_title(axis, t, size=12):
    axis.has_title = True
    tf = axis.axis_title.text_frame
    tf.text = t
    r = tf.paragraphs[0].runs[0]
    r.font.size = Pt(size)
    r.font.bold = False
    r.font.color.rgb = MUTED


# ---------------------------------------------------------------------------
# Data for native charts
# ---------------------------------------------------------------------------
def osnr_vs_severity():
    """GSNR of the nominal link when a single span gets extra loss."""
    d = LinkDesign.nominal()
    sev = np.arange(0, 14.01, 0.5)
    curves = {}
    for span in (1, 10, 19):
        curves[span] = [snapshot(d, [Fault("fiber_degradation", span - 1, float(v))], noise=False)[1] if v
                        else snapshot(d, [], noise=False)[1] for v in sev]
    return sev, curves


# ---------------------------------------------------------------------------
def build():
    r = json.loads((C.RESULTS / "metrics.json").read_text())
    pm, cv, ext = r["paper_models"], r["cross_validation"], r["extensions"]
    loc, fl, ew, ds = r["localization"], r["flapping"], r["early_warning"], r["dataset"]
    lr, gda, svm = pm["Logistic Regression"], pm["GDA"], pm["Linear SVM"]
    agree = r["sklearn_validation"]
    pct = lambda v: f"{100 * v:.1f}%"

    # 1 — Title ---------------------------------------------------------------
    s = new_slide(dark=True)
    text(s, 0.8, 0.75, 9, 0.4, "UE24CS352A  ·  MACHINE LEARNING  ·  MINI-PROJECT", size=13, color=ICE, bold=True)
    text(s, 0.8, 1.45, 11.6, 2.2, "Link Failure Prediction & Localization in Cloud-Scale Networks",
         size=46, font=HEAD, bold=True, color=WHITE)
    text(s, 0.8, 3.65, 11.5, 0.6, "Supervised learning on optical-line telemetry — "
         "Logistic Regression · GDA · Linear SVM, built from scratch", size=19, color=ICE)
    y0 = 5.0
    for i in range(10):                           # motif: the optical line
        x0 = 0.8 + i * 1.2
        shape(s, MSO_SHAPE.OVAL, x0, y0, 0.34, 0.34, fill=ORANGE if i == 6 else AQUA)
        if i < 9:
            line(s, x0 + 0.34, y0 + 0.17, x0 + 1.2, y0 + 0.17, color=ICE, w=2)
    text(s, 0.8, 5.5, 11.5, 0.4, "Each dot is an amplifier site on one link. One of them is failing. "
         "Can a model spot it from telemetry alone, and say which one?", size=13, color=ICE)
    text(s, 0.8, 6.35, 5.5, 0.35, SANJAY, size=17, bold=True, color=WHITE)
    text(s, 0.8, 6.7, 5.5, 0.3, "PES1UG24CS194", size=12, color=ICE)
    text(s, 4.3, 6.35, 5.5, 0.35, KANAK, size=17, bold=True, color=WHITE)
    text(s, 4.3, 6.7, 5.5, 0.3, "PES1UG24CS212", size=12, color=ICE)
    notes(s, SANJAY, "Introduce the team. One sentence: we reproduce and extend a Stanford paper that predicts "
          "optical link failures from span-loss and amplifier-gain telemetry and then localises the failing span. "
          "Split: Sanjay covers problem, data and preprocessing, localisation and demo; Kanak covers models, "
          "results, implementation and conclusions.")

    # 2 — Problem -----------------------------------------------------------------
    s = new_slide("Problem", SANJAY)
    title(s, "Find failing links before customers do",
          "Every end-to-end link crosses 19 fibre spans and amplifiers — any one of them can degrade")
    bullets(s, 0.6, 1.95, 5.3, 3.2, [
        [("Cloud traffic rides on ", {}), ("IP-over-optical", {"bold": True}), (" links hundreds of km long", {})],
        [("Fibres age, bend or get cut; amplifiers lose gain; connectors ", {}), ("flap", {"bold": True})],
        [("Each fault erodes the ", {}), ("OSNR", {"bold": True}), (" margin until traffic drops", {})],
        [("Goal: ", {"bold": True, "color": ORANGE}), ("flag marginal links ", {"bold": True}),
         ("before", {"bold": True, "color": ORANGE}), (" the outage and ", {}), ("say which span", {"bold": True})],
    ], size=17)
    flows = [("Today: reactive", ORANGE, TINT_O, ["Degrades", "Outage", "Ticket", "Search", "Repair"], 1),
             ("With this project: proactive", AQUA, TINT_A, ["Degrades", "ML alarm", "Span found",
                                                             "Repair", "No outage"], None)]
    for k, (head, col, tint, steps, bad) in enumerate(flows):
        y = 2.0 + k * 2.35
        card(s, 6.3, y, 6.45, 2.05, tint)
        text(s, 6.55, y + 0.18, 6, 0.4, head, size=16, bold=True, color=col)
        for i, st in enumerate(steps):
            x = 6.55 + i * 1.22
            c = shape(s, MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, x, y + 0.8, 1.3, 0.8,
                      fill=ORANGE if (bad is not None and i == bad) else (col if i == len(steps) - 1 else WHITE),
                      line=col, line_w=1.25)
            dark_fill = (bad is not None and i == bad) or i == len(steps) - 1
            label_in(c, st, size=11, color=WHITE if dark_fill else INK, bold=dark_fill)
    text(s, 0.6, 5.6, 5.3, 1.1, [[("Paper's claim: ", {"bold": True, "color": NAVY}),
                                  ("two aggregated telemetry features and a linear classifier are enough to "
                                   "tell healthy links from failing ones.", {"italic": True})]], size=15)
    notes(s, SANJAY, "Contrast the two flows. Reactive: customers notice first, then engineers search 19 spans. "
          "Proactive: the model raises the alarm during the slow degradation and points at the span, so the repair "
          "can be scheduled. We test the paper's claim that this needs only two features and a linear model.")

    # 3 — Telemetry / system diagram ----------------------------------------------------
    s = new_slide("Problem", SANJAY)
    title(s, "Where the data comes from", "One E2E link, as in Fig. 1 of the paper — every 2 minutes each span reports 4 numbers")
    y = 2.3
    def node(x, w, label, fill, size=12):
        b = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.9, fill=fill, radius=0.12)
        label_in(b, label, size=size)
    node(0.6, 0.9, "Tx", NAVY, 15)
    line(s, 1.5, y + 0.45, 1.7, y + 0.45, color=MUTED)
    node(1.7, 0.95, "ROADM", BLUE, 11)
    x = 2.65
    for i, span_no in enumerate([1, 2, 3, 19]):
        fail = span_no == 3
        if span_no == 19:
            text(s, x + 0.02, y + 0.15, 0.35, 0.5, "…", size=22, color=MUTED, bold=True)
            x += 0.35
        line(s, x, y + 0.45, x + 0.2, y + 0.45, color=MUTED)
        fib = shape(s, MSO_SHAPE.FLOWCHART_MAGNETIC_DISK, x + 0.2, y + 0.15, 0.85, 0.6,
                    fill=TINT_O if fail else TINT, line=ORANGE if fail else BLUE, line_w=1.5)
        label_in(fib, f"span {span_no}", size=10, color=ORANGE if fail else NAVY)
        line(s, x + 1.05, y + 0.45, x + 1.2, y + 0.45, color=MUTED)
        shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, x + 1.2, y + 0.18, 0.55, 0.55,
              fill=ORANGE if fail else AQUA, rot=90)
        x += 1.75
    line(s, x, y + 0.45, x + 0.2, y + 0.45, color=MUTED)
    node(x + 0.2, 0.95, "ROADM", BLUE, 11)
    line(s, x + 1.15, y + 0.45, x + 1.35, y + 0.45, color=MUTED)
    node(x + 1.35, 0.9, "Rx", NAVY, 15)
    text(s, 0.6, y + 1.12, 12.1, 0.4, "fibre span  →  amplifier (gives back the power the span lost)  "
         "× 19 spans  →  receiver.   Orange = the failing span.", size=12, color=MUTED)
    metrics_ = [("Span loss", "dB measured on the fibre", ORANGE), ("Target span loss", "planned value", MUTED),
                ("Amplifier gain", "dB the amplifier adds", AQUA), ("Target gain", "planned value", MUTED)]
    for i, (a, b, col) in enumerate(metrics_):
        xx = 0.6 + i * 3.08
        card(s, xx, 4.0, 2.85, 1.15)
        badge(s, xx + 0.2, 4.3, i + 1, d=0.5, fill=col if col != MUTED else BLUE, size=14)
        text(s, xx + 0.85, 4.25, 1.95, 0.4, a, size=15, bold=True, color=NAVY)
        text(s, xx + 0.85, 4.62, 1.95, 0.4, b, size=11, color=MUTED)
    card(s, 0.6, 5.45, 12.13, 1.3, NAVY)
    text(s, 0.95, 5.62, 11.5, 1.0,
         [[("Healthy link: ", {"bold": True, "color": WHITE}),
           ("every span loses what was planned and every amplifier gives it back, so all diffs ≈ 0.", {})],
          [("Failing link: ", {"bold": True, "color": ORANGE}),
           ("one span loses more, or one amplifier gives less, and the signal arrives weaker and noisier.", {})]],
         size=15, color=ICE)
    notes(s, SANJAY, "Walk the signal left to right. Each amplifier is designed to give back exactly what the "
          "previous fibre span lost. The four telemetry numbers are exactly the paper's raw features. "
          "The orange span is the kind of fault we want to find.")

    # 4 — Paper vs us -----------------------------------------------------------------------
    s = new_slide("Problem", SANJAY)
    title(s, "What the paper did, and what we added", "Bakhtiari, \"Link Failure Prediction and Localization in "
          "Cloud Scale Networks using Supervised Learning\", Stanford")
    rows = [("Data", "Live operator data (private)", "Physics-based simulator, same format, ground-truth fault spans"),
            ("Features", "X1 = Σ gain diff,  X2 = Σ loss diff", "Same, plus 38 per-span residuals (extension)"),
            ("Models", "LR (Newton), GDA, linear SVM", "Same three, written from scratch + RBF SVM"),
            ("Evaluation", "Test accuracy (75/25)", "Accuracy, P/R/F1, ROC, 5-fold CV, unseen-route CV"),
            ("Localisation", "Claimed, not measured", "Measured: span + element, flapping, early warning"),
            ("Deliverable", "Report", "Pipeline, tests, live demo app, CLI")]
    x0s, ws = [0.6, 2.75, 7.0], [2.0, 4.1, 5.73]
    for j, h in enumerate(["", "Paper", "This project"]):
        if h:
            c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x0s[j], 1.85, ws[j], 0.5,
                      fill=GREY if j == 1 else BLUE, radius=0.2)
            label_in(c, h, size=14)
    for i, (a, b, c_) in enumerate(rows):
        yy = 2.5 + i * 0.7
        if i % 2 == 0:
            card(s, 0.6, yy - 0.06, 12.13, 0.62, TINT, radius=0.12)
        text(s, 0.8, yy + 0.06, 1.9, 0.45, a, size=15, bold=True, color=NAVY)
        text(s, 2.95, yy + 0.07, 3.9, 0.45, b, size=14, color=MUTED)
        text(s, 7.2, yy + 0.07, 5.4, 0.45, c_, size=14, color=INK)
    notes(s, SANJAY, "Be upfront: the paper's data is private and its appendix link is dead, so we built a "
          "simulator with the same format. That gave us something the paper never had: the true location of every "
          "fault, so we could actually measure localisation.")

    # 5 — Dataset -------------------------------------------------------------------------------
    s = new_slide("Data", SANJAY)
    title(s, "Dataset: same shape as the paper", "Telemetry simulator: ASE + nonlinear-noise optical link budget, realistic faults and gaps")
    stats = [(str(ds["links"]), "directional E2E links (10 routes × 2)"), (str(C.N_SPANS), "spans per link"),
             (f"{ds['samples']:,}", "samples, one per link per 2-min bin"), (pct(ds["not_healthy_fraction"]), "labelled not healthy")]
    for i, (v, l) in enumerate(stats):
        stat(s, 0.6 + i * 3.08, 1.85, 2.85, v, l, color=ORANGE if i == 3 else BLUE)
    ft = ds["fault_type_counts"]
    total = sum(v for k, v in ft.items() if k != "none")
    text(s, 0.6, 3.5, 6, 0.4, "Injected faults (bins)", size=16, bold=True, color=NAVY)
    faults = [("Fibre degradation", "1–14 dB, slow ramp or step", ft["fiber_degradation"], ORANGE),
              ("Amplifier degradation", "gain ↓, noise figure ↑", ft["amplifier_degradation"], BLUE),
              ("Flapping connector", "intermittent loss", ft["flapping"], AQUA),
              ("Fibre cut", "+40 dB, no reading (LOS)", ft["fiber_cut"], NAVY)]
    for i, (a, b, nbin, col) in enumerate(faults):
        yy = 4.0 + i * 0.68
        text(s, 0.6, yy, 2.6, 0.3, a, size=14, bold=True)
        text(s, 0.6, yy + 0.29, 2.6, 0.3, b, size=11, color=MUTED)
        bw = 2.9 * nbin / max(ft["fiber_degradation"], 1)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 3.25, yy + 0.08, bw, 0.36, fill=col, radius=0.3)
        text(s, 3.35 + bw, yy + 0.1, 0.9, 0.35, f"{nbin:,}", size=13, color=INK, bold=True)
    text(s, 0.6, 6.72, 6, 0.3, f"{total:,} faulty bins in total; the rest are healthy operation.",
         size=11, color=MUTED)
    card(s, 7.0, 3.5, 5.73, 3.3)
    text(s, 7.3, 3.68, 5.2, 0.4, "Labels from OSNR, as in the paper", size=16, bold=True, color=NAVY)
    bullets(s, 7.3, 4.2, 5.2, 2.6, [
        [("Not healthy", {"bold": True, "color": ORANGE}), (" if OSNR < 17 dB, or Rx power > 7 dB low, or loss of signal", {})],
        "Noise: 0.12 dB on loss, 0.06 dB on gain",
        "~2 % cells missing, whole bins missing, no data during cuts",
        "75 / 25 split → 4500 train, 1500 test",
    ], size=14, gap=7)
    notes(s, SANJAY, "Q&A: why simulate? The paper's data is private. The simulator uses constant-gain amplifiers "
          "(a power deficit propagates downstream), ASE OSNR per amplifier = 58 + P_in − NF, and the GN model for "
          "nonlinear noise. Labels are the physical outcome, not a rule on X1/X2, so the classifier has to learn it.")

    # 6 — Physics: why position matters --------------------------------------------------------------
    s = new_slide("Data", SANJAY)
    title(s, "The physics behind the labels", "Same extra loss, different damage: it depends on where in the line it happens")
    sev, curves = osnr_vs_severity()
    cd = XyChartData()
    cols = {1: ORANGE, 10: BLUE, 19: AQUA}
    for span, ys in curves.items():
        ser = cd.add_series(f"Fault in span {span}")
        for xv, yv in zip(sev, ys):
            ser.add_data_point(float(xv), round(float(yv), 2))
    thr = cd.add_series("Unhealthy below 17 dB")
    thr.add_data_point(0.0, 17.0)
    thr.add_data_point(14.0, 17.0)
    rxl = cd.add_series("Rx-power alarm (7 dB)")
    rxl.add_data_point(7.0, 10.0)
    rxl.add_data_point(7.0, 21.0)
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(0.5), Inches(1.75),
                            Inches(7.6), Inches(5.05), cd)
    ch = gf.chart
    style_chart(ch)
    ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 10, 21
    ch.category_axis.minimum_scale, ch.category_axis.maximum_scale = 0, 14
    ch.category_axis.major_unit = 2
    ch.category_axis.has_major_gridlines = False
    axis_title(ch.value_axis, "Link OSNR [dB]")
    axis_title(ch.category_axis, "Extra loss injected into one span [dB]")
    for ser, col in zip(ch.plots[0].series, [ORANGE, BLUE, AQUA, MUTED, GREY]):
        ser.smooth = False
        ser.format.line.color.rgb = col
        ser.format.line.width = Pt(3 if col in (ORANGE, BLUE, AQUA) else 1.75)
        if col in (MUTED, GREY):
            ser.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    c1, c19 = curves[1], curves[19]
    i7 = int(np.where(sev == 7.0)[0][0])
    card(s, 8.5, 1.9, 4.23, 2.25, TINT_O)
    text(s, 8.75, 2.05, 3.8, 0.4, "7 dB in span 1", size=16, bold=True, color=ORANGE)
    text(s, 8.75, 2.5, 3.8, 1.6, f"OSNR {c1[i7]:.1f} dB. Every one of the 19 downstream amplifiers now works "
         "with a weaker signal and adds proportionally more noise.", size=14)
    card(s, 8.5, 4.35, 4.23, 2.25, TINT_A)
    text(s, 8.75, 4.5, 3.8, 0.4, "7 dB in span 19", size=16, bold=True, color=AQUA)
    text(s, 8.75, 4.95, 3.8, 1.6, f"OSNR {c19[i7]:.1f} dB. Only the last amplifier is affected — but the "
         "receiver is now starved of power, so the Rx-power alarm fires.", size=14)
    notes(s, SANJAY, "This chart is computed live from our simulator. Key message for later: X1 and X2 are sums "
          "over spans, so they cannot see WHERE the loss happened. That position effect is the main reason "
          "accuracy is ~97 % rather than 99 %.")

    # 7 — Preprocessing ----------------------------------------------------------------------------------
    s = new_slide("Method", SANJAY)
    title(s, "Preprocessing & feature engineering", "Paper §III: from 4 × 19 raw values to 2 features per link per bin")
    steps = [("Fibre-cut fill", "LOS-flagged gaps ← target + 40 dB (the paper's rule)"),
             ("Interpolate", "linear in time, per link direction and per span"),
             ("Span diffs (Eq. 1)", "x1 = gain − target gain\nx2 = target loss − span loss"),
             ("Aggregate (Eq. 2)", "X1 = Σ x1 ,  X2 = Σ x2\nover the 19 spans"),
             ("Join labels", "OSNR-based y ∈ {0, 1}")]
    for i, (a, b) in enumerate(steps):
        x0 = 0.6 + i * 2.5
        card(s, x0, 1.95, 2.2, 2.45)
        badge(s, x0 + 0.2, 2.15, i + 1)
        text(s, x0 + 0.2, 2.85, 1.85, 0.45, a, size=15, bold=True, color=NAVY)
        text(s, x0 + 0.2, 3.3, 1.85, 1.0, b, size=12)
        if i < 4:
            arrow(s, x0 + 2.22, 3.0, 0.26, 0.3)
    card(s, 0.6, 4.75, 6.0, 2.0, NAVY)
    text(s, 0.9, 4.92, 5.5, 0.4, "Why these two features?", size=16, bold=True, color=WHITE)
    text(s, 0.9, 5.37, 5.5, 1.3, "X1 + X2 is the net optical power the link has lost versus its design. "
         "They don't depend on span lengths, so one model serves every link.", size=14, color=ICE)
    card(s, 6.9, 4.75, 5.83, 2.0, TINT)
    text(s, 7.2, 4.92, 5.3, 0.4, "Data quality handled", size=16, bold=True, color=NAVY)
    tele_missing = "4,862"
    bullets(s, 7.2, 5.37, 5.3, 1.3, [f"{tele_missing} missing cells filled (≈ 2 % of telemetry)",
                                     "cuts filled, not interpolated: no light means no reading"], size=13, gap=5)
    notes(s, SANJAY, "Walk through linkfail/preprocessing.py. The fibre-cut special case comes straight from the "
          "paper: a gap caused by a root-cause failure is filled with the expected cut value, not interpolated. "
          "Handover to Kanak for the models.")

    # 8 — Models ----------------------------------------------------------------------------------------
    s = new_slide("Method", KANAK)
    title(s, "Three linear classifiers, written from scratch", "NumPy implementations in linkfail/models.py, checked against scikit-learn")
    cards = [("Logistic Regression", BLUE, "DISCRIMINATIVE",
              ["J(θ) = −1/n Σ [y log h + (1−y) log(1−h)]", "Newton's method: θ ← θ − H⁻¹∇J",
               f"Converges in {lr['newton_iterations']} iterations, J = {lr['final_cost']:.4f}"]),
             ("GDA", ORANGE, "GENERATIVE",
              ["y ~ Bernoulli(φ),  x | y ~ N(μ_y, Σ)", "Closed-form MLE: φ, μ₀, μ₁, shared Σ",
               "Shared Σ ⇒ straight boundary at p(y|x) = 0.5"]),
             ("Linear SVM", AQUA, "MAX-MARGIN",
              ["min λ/2‖w‖² + 1/n Σ max(0, 1 − y(wᵀx + b))", "Mini-batch Pegasos sub-gradient solver",
               "Iterate averaging, 200 epochs"])]
    for i, (name, col, kind, lines) in enumerate(cards):
        x0 = 0.6 + i * 4.1
        card(s, x0, 1.85, 3.85, 4.15)
        shape(s, MSO_SHAPE.OVAL, x0 + 0.3, 2.1, 0.4, 0.4, fill=col)
        text(s, x0 + 0.85, 2.08, 2.9, 0.45, name, size=19, bold=True, color=NAVY, font=HEAD)
        text(s, x0 + 0.3, 2.68, 3.3, 0.3, kind, size=11, bold=True, color=MUTED)
        bullets(s, x0 + 0.3, 3.1, 3.35, 2.3, lines, size=13, marker_color=col)
        pill = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x0 + 0.3, 5.3, 3.25, 0.45, fill=WHITE, line=col, radius=0.5)
        label_in(pill, f"matches scikit-learn on {pct(agree[name]['prediction_agreement'])} of test points",
                 size=11, color=col)
    text(s, 0.6, 6.3, 12.1, 0.5, "All three take raw (X1, X2). LR and SVM standardise internally and report θ in "
         "raw dB units, so coefficients are directly interpretable.", size=13, color=MUTED)
    notes(s, KANAK, "Why Newton? Two features means a 3×3 Hessian, so each step is trivial and convergence is "
          "quadratic (11 steps). GDA vs LR: GDA assumes Gaussian classes, which helps when true and hurts when not. "
          "Pegasos: projected stochastic sub-gradient on the primal, step size 1/(λt). Validation against "
          "scikit-learn proves the implementations are correct.")

    # 9 — Data view -------------------------------------------------------------------------------------
    s = new_slide("Results", KANAK)
    title(s, "What the model sees", "Each dot is one link in one 2-minute bin, coloured by its OSNR label")
    figure(s, "dataset_scatter", 0.45, 1.8, w=8.5)
    items = [(BLUE, "Healthy", "cluster tightly at (0, 0)"),
             (ORANGE, "Amplifier faults", "spread left: X1 ≪ 0"),
             (ORANGE, "Fibre faults", "spread down: X2 ≪ 0"),
             (NAVY, "Fibre cuts", "a separate island at X2 ≈ −40 dB"),
             (AQUA, "Boundary", "nearly a straight line, with overlap at the edge")]
    for i, (col, a, b) in enumerate(items):
        yy = 2.0 + i * 0.92
        shape(s, MSO_SHAPE.OVAL, 9.25, yy + 0.08, 0.26, 0.26, fill=col)
        text(s, 9.7, yy, 3.0, 0.35, a, size=15, bold=True)
        text(s, 9.7, yy + 0.35, 3.0, 0.4, b, size=13, color=MUTED)
    notes(s, KANAK, "The geometry explains everything that follows: a near-linear boundary (so linear models "
          "work), and a multi-cluster failure class (so GDA's single-Gaussian assumption will struggle).")

    # 10 — Decision boundaries ------------------------------------------------------------------------------
    s = new_slide("Results", KANAK)
    title(s, "Decision boundaries on the held-out test set", "1500 samples never used in training — compare with Fig. 4–5 of the paper")
    figure(s, "decision_boundaries", 0.45, 1.75, w=12.4)
    for i, (name, m, col) in enumerate([("Logistic Regression", lr, BLUE), ("GDA", gda, ORANGE), ("Linear SVM", svm, AQUA)]):
        x0 = 0.6 + i * 4.12
        t = m["theta"]
        stat(s, x0, 5.5, 3.88, pct(m["accuracy"]),
             f"{name}  ·  θ = ({t[0]:.2f}, {t[1]:.2f}, {t[2]:.2f})", color=col, vsize=28, h=1.3)
    notes(s, KANAK, "LR and SVM learn nearly identical lines with θ1 ≈ θ2: the rule is 'alarm when X1 + X2 drops "
          "below about −7 dB', i.e. net power deficit, which is physically right. GDA's line is tilted by the "
          "fibre-cut cluster and misses many failures.")

    # 11 — Results vs paper (native chart) -------------------------------------------------------------------
    s = new_slide("Results", KANAK)
    title(s, "Results vs. the paper", "Same ranking; the gap has a physical explanation (next slides)")
    cd = CategoryChartData()
    cd.categories = ["Logistic Regression", "GDA", "Linear SVM"]
    cd.add_series("Paper", (99.0, 97.0, 99.0))
    cd.add_series("This project", tuple(round(100 * pm[n]["accuracy"], 1) for n in ("Logistic Regression", "GDA", "Linear SVM")))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.75),
                            Inches(7.4), Inches(5.0), cd)
    ch = gf.chart
    style_chart(ch, size=13)
    ch.value_axis.minimum_scale, ch.value_axis.maximum_scale, ch.value_axis.major_unit = 80, 100, 5
    axis_title(ch.value_axis, "Test accuracy [%]  (axis starts at 80)")
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 70, -8
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = '0.0"%"', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold = Pt(13), True
    for ser, col in zip(plot.series, [GREY, BLUE]):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
    stat(s, 8.3, 1.9, 4.43, f"{lr['roc_auc']:.3f}",
         "ROC-AUC of LR and linear SVM: ranking is near-perfect, errors sit at the boundary", color=BLUE)
    stat(s, 8.3, 3.45, 4.43, pct(cv["Linear SVM"]["route_grouped_5fold_mean"]),
         "on routes never seen in training (GroupKFold): no overfitting to specific links", color=AQUA)
    stat(s, 8.3, 5.0, 4.43, f"{gda['recall']:.2f}", "GDA recall: it misses 4 in 10 failures despite high precision",
         color=ORANGE)
    notes(s, KANAK, "Grey bars are the paper's numbers, blue are ours. The ordering LR ≈ SVM > GDA is reproduced. "
          "Our absolute numbers are lower because our labels come from OSNR physics in which fault position matters.")

    # 12 — Robustness ------------------------------------------------------------------------------------------
    s = new_slide("Results", KANAK)
    title(s, "Beyond accuracy: robust on unseen links", "Precision/recall on the test set and cross-validation over all 6000 samples")
    hdr = ["Classifier", "Acc.", "Precision", "Recall", "F1", "AUC", "5-fold CV", "Unseen routes"]
    data = [hdr]
    for name in ("Logistic Regression", "GDA", "Linear SVM"):
        m, c = pm[name], cv[name]
        data.append([name, pct(m["accuracy"]), f"{m['precision']:.3f}", f"{m['recall']:.3f}", f"{m['f1']:.3f}",
                     f"{m['roc_auc']:.3f}", f"{pct(c['stratified_5fold_mean'])} ± {100 * c['stratified_5fold_std']:.1f}",
                     f"{pct(c['route_grouped_5fold_mean'])} ± {100 * c['route_grouped_5fold_std']:.1f}"])
    widths = [2.6, 1.1, 1.25, 1.1, 1.0, 1.0, 1.9, 2.18]
    tbl = s.shapes.add_table(len(data), len(hdr), Inches(0.6), Inches(1.85), Inches(sum(widths)),
                             Inches(0.45 * len(data))).table
    for j, wdt in enumerate(widths):
        tbl.columns[j].width = Inches(wdt)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (TINT if i % 2 else WHITE)
            p = cell.text_frame.paragraphs[0]
            p.text = ""
            run = p.add_run()
            run.text = val
            run.font.size = Pt(14)
            run.font.name = BODY
            run.font.bold = i == 0 or j == 0 or (row[0] == "GDA" and j == 3)
            run.font.color.rgb = WHITE if i == 0 else (ORANGE if (row[0] == "GDA" and j == 3) else INK)
    figure(s, "confusion_matrices", 0.6, 3.95, w=9.0, h=2.85)
    figure(s, "roc_curves", 9.75, 3.9, w=3.0, h=2.95)
    notes(s, KANAK, "'Unseen routes' holds out whole fibre routes, both directions, so it's accuracy on a link the "
          "model has never seen; it matches the random split, so the model generalises. Confusion matrices: LR and "
          "SVM make ~26 errors each way; GDA misses 151 failures.")

    # 13 — Discussion ---------------------------------------------------------------------------------------------
    s = new_slide("Results", KANAK)
    title(s, "Why GDA trails — and why we're below 99 %", "The ranking matches the paper; both gaps have a clear cause")
    cd = CategoryChartData()
    names = ["Logistic Regression", "GDA", "Linear SVM", "SVM, RBF kernel *", "LR + 38 per-span features *"]
    vals = [pm["Logistic Regression"]["accuracy"], pm["GDA"]["accuracy"], pm["Linear SVM"]["accuracy"],
            ext["SVM (RBF kernel)"]["accuracy"], ext["LR + 38 per-span residuals"]["accuracy"]]
    cd.categories = list(reversed(names))
    cd.add_series("Test accuracy", tuple(round(100 * v, 1) for v in reversed(vals)))
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.45), Inches(1.75), Inches(6.9), Inches(4.7), cd)
    ch = gf.chart
    style_chart(ch, legend=False, size=12)
    ch.has_title = False
    ch.value_axis.minimum_scale, ch.value_axis.maximum_scale, ch.value_axis.major_unit = 85, 100, 5
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    plot.data_labels.number_format, plot.data_labels.number_format_is_linked = '0.0"%"', False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size, plot.data_labels.font.bold = Pt(12), True
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = BLUE
    for idx, nm in enumerate(reversed(names)):
        pt = ser.points[idx]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = AQUA if nm.endswith("*") else (ORANGE if nm == "GDA" else BLUE)
    text(s, 0.6, 6.5, 6.6, 0.35, "* extensions beyond the paper  ·  axis starts at 85 %", size=11, color=MUTED)
    card(s, 7.7, 1.85, 5.03, 2.3, TINT_O)
    text(s, 7.95, 2.0, 4.6, 0.4, "GDA's assumption breaks", size=17, bold=True, color=ORANGE)
    text(s, 7.95, 2.45, 4.6, 1.65, "'Not healthy' is three clusters (amplifier, fibre, cut), not one Gaussian. "
         f"They stretch the shared Σ and tilt the boundary: recall falls to {gda['recall']:.2f}.", size=14)
    card(s, 7.7, 4.35, 5.03, 2.4, TINT)
    text(s, 7.95, 4.5, 4.6, 0.4, "Σ over spans hides where", size=17, bold=True, color=BLUE)
    text(s, 7.95, 4.95, 4.6, 1.75, f"A nonlinear RBF kernel adds just {100 * (ext['SVM (RBF kernel)']['accuracy'] - svm['accuracy']):.1f} "
         f"pts; giving the model per-span features adds more. The limit is information lost in the sum, "
         "not the linear model.", size=14)
    notes(s, KANAK, "Key Q&A point. If the problem were non-linearity, the RBF SVM would jump; it barely moves. "
          "Adding per-span residuals helps more, confirming the missing ingredient is fault position, which the "
          "paper's two summed features discard by design. The paper's own conclusion — linear is enough — holds.")

    # 14 — Localisation ------------------------------------------------------------------------------------------------
    s = new_slide("Operations", SANJAY)
    title(s, "From alarm to work order", "Localisation, flapping detection and early warning")
    figure(s, "localization_example", 0.45, 1.75, w=7.7)
    text(s, 0.6, 4.75, 7.5, 1.6,
         [[("Localise: ", {"bold": True, "color": NAVY}), ("the span with the largest positive residual. Loss "
                                                         "residual → fibre, gain residual → amplifier, ≥ 25 dB → cut.", {})],
          [("Flapping: ", {"bold": True, "color": NAVY}), ("≥ 4 health toggles within 30 minutes.", {})],
          [("Early warning: ", {"bold": True, "color": NAVY}), ("first alarm vs. first outage bin, per slow degradation.", {})]],
         size=14)
    stat(s, 8.5, 1.8, 4.23, pct(loc["span_accuracy"]), f"span + element correct on detected failures (n = {loc['n_evaluated']})",
         color=AQUA)
    stat(s, 8.5, 3.3, 4.23, f"{fl['events_detected']}/{fl['flapping_events']}",
         f"flapping events caught ({pct(fl['false_alarm_rate_non_flapping_bins'])} false alarms on other bins)")
    stat(s, 8.5, 4.8, 4.23, f"{ew['flagged_before_or_at_outage']}/{ew['events_with_outage']}",
         f"outages flagged in time; {ew['flagged_strictly_before_outage']} strictly before (mean {ew['mean_lead_minutes']:.0f} min ahead)",
         color=ORANGE)
    text(s, 0.6, 6.45, 12, 0.4, f"Harder test, every fault incl. minor ones: {pct(loc['span_accuracy_all_fault_samples'])} "
         f"span accuracy  ·  learnt multinomial-LR localiser: {pct(loc['ml_localizer_span_accuracy'])}", size=12, color=MUTED)
    notes(s, SANJAY, "The paper only claimed localisation; we measured it because the simulator knows the true "
          "span. 100 % on detected failures is expected: a real failure leaves a residual far above the 0.1 dB "
          "noise. The honest harder number, including minor faults, is 91.6 %.")

    # 15 — Live demo ------------------------------------------------------------------------------------------------------
    s = new_slide("Demo", SANJAY)
    title(s, "Live demo", "streamlit run app.py  —  inject a fault, watch three models and the localiser react")
    picture(s, ASSETS / "app_fault_injection.png", 5.55, 1.7, w=7.18, h=5.15, border=True)
    demo = [("Healthy baseline", "?fault=none", "all three say healthy"),
            ("9 dB fibre loss, span 8", "?fault=fiber_degradation&span=8&sev=9", "LR & SVM alarm, GDA misses it, span 8 localised"),
            ("Same loss, span 19", "&span=19", "OSNR barely moves — the position effect"),
            ("Fibre cut", "?fault=fiber_cut&span=12", "everyone alarms, localised as a cut"),
            ("Network replay", "tab 2, link R07-ZA", "10 h timeline with flapping alarms")]
    for i, (a, q, b) in enumerate(demo):
        yy = 1.8 + i * 1.0
        badge(s, 0.6, yy + 0.05, i + 1, d=0.48, size=14)
        text(s, 1.25, yy, 4.1, 0.35, a, size=15, bold=True, color=NAVY)
        text(s, 1.25, yy + 0.33, 4.1, 0.3, b, size=12)
        text(s, 1.25, yy + 0.6, 4.1, 0.3, q, size=10, color=MUTED, font=MONO)
    notes(s, SANJAY, "Scenario URLs append to http://localhost:8501 . The screenshot shows scenario 2: X2 = −8.4 dB, "
          "OSNR 15.4 dB, LR and SVM flag it, GDA says healthy (wrong), localiser points to span 8 fibre. "
          "Fallback if the app fails: python predict.py --snapshot examples/degraded_span7.csv")

    # 16 — Implementation ------------------------------------------------------------------------------------------------
    s = new_slide("Engineering", KANAK)
    title(s, "Implementation & repository", "github.com/jsanjayvarma06-oss/ML-Project  ·  one command reproduces every number here")
    card(s, 0.6, 1.85, 6.1, 4.55, NAVY, radius=0.04)
    tree = ["run_pipeline.py      one-command experiment", "app.py               Streamlit live demo",
            "predict.py           CLI inference", "linkfail/", "  simulator.py       telemetry + link budget",
            "  preprocessing.py   Eq. 1 / Eq. 2 features", "  models.py          LR · GDA · SVM",
            "  localization.py    span + flapping", "  evaluation.py      metrics, lead time",
            "  pipeline.py        CV, extensions", "tests/               12 pytest tests",
            "docs/                write-up + slides"]
    text(s, 0.9, 2.15, 5.6, 4.1, tree, size=14, font=MONO, color=WHITE, space_after=3)
    pts = [("Reproducible", "fixed seed; whole pipeline in ~20 s; reruns give identical metrics"),
           ("Validated", "99.8–100 % prediction agreement with scikit-learn"),
           ("Tested", "physics, interpolation, Eq. 2 identity, localisation"),
           ("Honest evaluation", "held-out test, 5-fold CV, unseen-route CV, out-of-fold timelines")]
    for i, (a, b) in enumerate(pts):
        yy = 1.95 + i * 1.12
        badge(s, 7.1, yy, "✓", d=0.48, fill=AQUA, size=14)
        text(s, 7.75, yy - 0.02, 5.0, 0.35, a, size=16, bold=True, color=NAVY)
        text(s, 7.75, yy + 0.35, 5.0, 0.6, b, size=13)
    text(s, 7.1, 6.45, 5.6, 0.4, f"Contributions — {SANJAY}: simulator, preprocessing, localisation, demo  ·  "
         f"{KANAK}: models, evaluation, extensions, report", size=11, color=MUTED)
    notes(s, KANAK, "Show the repo briefly if time allows. Emphasise reproducibility: same seed, same numbers. "
          "Tests check the models against scikit-learn and the Eq. 2 identity (features equal the sum of residuals).")

    # 17 — Conclusions -----------------------------------------------------------------------------------------------
    s = new_slide("Wrap-up", KANAK, dark=True)
    title(s, "Conclusions", dark=True)
    pts = [("Two telemetry features + a linear model", f"detect failing links at {pct(svm['accuracy'])}, equally "
                                                      "well on links never seen in training"),
           ("The paper's ranking reproduced", "LR ≈ linear SVM > GDA; GDA's Gaussian assumption fails on "
                                              "multi-cluster failures"),
           ("Localisation comes almost free", "per-span residuals pinpoint the span and the element: fibre, amplifier or cut"),
           ("Next", "forecast OSNR for longer lead time · position-aware features · validate on real operator data")]
    for i, (a, b) in enumerate(pts):
        yy = 1.55 + i * 1.25
        badge(s, 0.7, yy, i + 1, fill=AQUA if i < 3 else ORANGE)
        text(s, 1.5, yy - 0.04, 11, 0.45, a, size=21, bold=True, color=WHITE, font=HEAD)
        text(s, 1.5, yy + 0.46, 11, 0.5, b, size=15, color=ICE)
    text(s, 0.7, 6.55, 11.5, 0.4, "Thank you — questions?", size=20, bold=True, color=WHITE)
    notes(s, KANAK, "Close on the one-line takeaway: cheap features plus a linear model catch failing links, and "
          "the same telemetry tells you where to send the engineer. Then open for questions; backup slides follow.")

    # 18 — Backup: Q&A -------------------------------------------------------------------------------------------------
    s = new_slide("Backup", None)
    title(s, "Backup · likely questions", "Short answers, with the evidence behind them")
    qa = [("Why synthetic data?", "Paper's data is private; its link is dead. Same format, plus true fault spans — "
                                  "which let us measure localisation."),
          ("Isn't 96 % worse than 99 %?", "Our labels follow OSNR physics where fault position matters; "
                                          "the RBF SVM barely helps, so the linear model isn't the limit."),
          ("Any data leakage?", "Unseen-route CV holds out whole routes: 96.9 %, same as the random split."),
          ("Why Newton, not gradient descent?", "3×3 Hessian is free to invert; quadratic convergence in 11 steps."),
          ("Why does GDA lose?", "Failures form 3 clusters; one Gaussian + shared Σ tilts the boundary. Recall 0.59."),
          ("Class imbalance?", "24.5 % positive; we report precision, recall, F1 and AUC, not only accuracy."),
          ("How are labels made?", "OSNR < 17 dB, Rx power > 7 dB low, or LOS — from an ASE + GN-model budget."),
          ("Real-world deployment?", "Features are sums of existing telemetry: O(19) per link per 2 min.")]
    for i, (q, a) in enumerate(qa):
        col, row = i % 2, i // 2
        x0, y0 = 0.6 + col * 6.12, 1.8 + row * 1.25
        card(s, x0, y0, 5.95, 1.12, TINT if (row + col) % 2 == 0 else WHITE)
        if (row + col) % 2:
            shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x0, y0, 5.95, 1.12, line=GRID, radius=0.06)
        text(s, x0 + 0.22, y0 + 0.1, 5.5, 0.35, q, size=14, bold=True, color=NAVY)
        text(s, x0 + 0.22, y0 + 0.45, 5.55, 0.65, a, size=12)
    notes(s, f"{SANJAY} & {KANAK}", "Keep this slide hidden during the talk and jump to it if a question comes up.")

    # 19 — Backup: equations & settings -------------------------------------------------------------------------
    s = new_slide("Backup", None)
    title(s, "Backup · equations & settings", "Everything needed to reproduce the models by hand")
    eq = [("Features (Eq. 1–2)", ["x1ₙ = Gₙ − G*ₙ ,   x2ₙ = L*ₙ − Lₙ", "X1 = Σₙ x1ₙ ,   X2 = Σₙ x2ₙ   (n = 1…19)"]),
          ("Logistic regression", ["h = σ(θᵀx),  ∇J = Xᵀ(h − y)/n + λθ", "H = XᵀSX/n + λI,  S = diag(h(1−h)),  λ = 10⁻⁴"]),
          ("GDA", ["φ = mean(y),  μₖ = mean(x | y = k)", "Σ = 1/n Σᵢ (xᵢ − μ_yᵢ)(xᵢ − μ_yᵢ)ᵀ,  w = Σ⁻¹(μ₁ − μ₀)"]),
          ("Linear SVM (Pegasos)", ["λ = 1/(C·n),  C = 1,  ηₜ = 1/(λt)", "batch 64, 200 epochs, average 2nd half"]),
          ("Labels / physics", ["OSNRᵢ = 58 + P_in,i − NFᵢ ;  SNR_NLI = −η − 2P", "unhealthy: OSNR < 17 dB ∨ Rx < −6 dBm ∨ LOS"]),
          ("Localisation", ["k* = argmaxₙ max(Lₙ − L*ₙ , G*ₙ − Gₙ)", "flapping: ≥ 4 toggles in 15 bins (30 min)"])]
    for i, (h, ls) in enumerate(eq):
        col, row = i % 2, i // 2
        x0, y0 = 0.6 + col * 6.12, 1.8 + row * 1.68
        card(s, x0, y0, 5.95, 1.5)
        text(s, x0 + 0.25, y0 + 0.13, 5.4, 0.35, h, size=14, bold=True, color=BLUE)
        text(s, x0 + 0.25, y0 + 0.55, 5.5, 0.9, ls, size=14, font="Cambria Math", space_after=4)
    notes(s, f"{SANJAY} & {KANAK}", "Reference slide for derivation questions.")

    prs.save(OUT)
    print(f"Wrote {OUT} ({_count[0]} slides)")


if __name__ == "__main__":
    build()
